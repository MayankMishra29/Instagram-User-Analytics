"""Generate the Trust Graph Tracker spreadsheet for the Implementation System."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "system" / "Trust-Graph-Tracker.xlsx"

INK = "1D2126"
ACCENT = "0F5E5A"
WARM = "C6613F"
TINT = "E5F0EE"
HEAD = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY = Font(name="Arial", size=10)
THIN = Side(style="thin", color="DFE2E6")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

STATES = "Active evaluation,Exploring,Promised follow-up,Invited revisit,Peer / referral,Dormant,Inactive,Do not contact,Won,Lost,Paused"
GATES = "G1 Reality,G2 Relevance,G3 Fit,G4 Safety,G5 Value,G6 Identity,Unknown"
POSTURES = "Curious,Skeptical,Guarded,Exploratory,Overwhelmed,Interested,Hesitant,Ready,Disengaged,Unknown"
TU = "P(outcome) weak,P(outcome) unknown,Value weak,Value unknown,Risk high,Risk unknown,Hard constraint,None - ready"
MOVES = "Mirror,Clarifier,Insight Gift,Story,Self-Disclosure,Direct answer,Offer (recap-reason-request),Rehydrate,Pause,Close the record"
RUNGS = "Attention,Curiosity,Interest,Exploration,Intent,Readiness"
ROUTES = "Existing relationship,Requested resource,Explicit public request,Introduction,Content bridge,Observation,Shared interest,Inbound,New contact"
CONF = "Directly stated,Strongly suggested,Plausible,Unknown"


def header(ws, cols, widths, fill=ACCENT):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = HEAD
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def dropdown(ws, col_letter, values, rows=500):
    dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
    dv.error = "Pick a value from the list"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{rows}")


def style_body(ws, ncols, rows=60):
    for r in range(2, rows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY
            cell.alignment = WRAP
            cell.border = BORDER


wb = Workbook()

# README
ws = wb.active
ws.title = "Start here"
ws.column_dimensions["A"].width = 110
lines = [
    ("THE NON-SALESY DM · TRUST GRAPH TRACKER", Font(bold=True, size=16, color=INK, name="Arial")),
    ("Part of the Implementation System for the Complete Edition. © Mayank Mishra / ThriveXLabs.", Font(italic=True, size=10, color="4A515B", name="Arial")),
    ("", BODY),
    ("HOW TO USE", Font(bold=True, size=12, color=ACCENT, name="Arial")),
    ("1. Trust Graph: one row per relationship. Update it during the RECORD step of the Daily OS (Chapter 14). Dropdowns keep states, gates and moves consistent.", BODY),
    ("2. Debugger Log: run the 13-question Trust Debugger (Appendix A) on significant, stalled or confusing threads only.", BODY),
    ("3. 30-Day Sprint: log each day of Chapter 15. A justified decision NOT to send counts as done.", BODY),
    ("4. Weekly Metrics: enter counts per cohort. Rates calculate automatically with explicit denominators. Always read counts next to percentages.", BODY),
    ("5. Post-Mortems: after any important thread (won, lost or stalled).", BODY),
    ("", BODY),
    ("RULES THAT KEEP THE TRACKER HONEST", Font(bold=True, size=12, color=ACCENT, name="Arial")),
    ("• Write the buyer's actual words in 'Last buyer signal'. Separate observation from interpretation.", BODY),
    ("• Never add a numeric trust score. Record evidence ('saw sample, asked about workload').", BODY),
    ("• 'Do not contact' overrides every reminder. Rows in that state are highlighted.", BODY),
    ("• Promises (mine) are checked first every day. Your reliability is part of your offer.", BODY),
    ("• Don't count follow-ups as new contacts. Track refusals separately from replies.", BODY),
    ("• Keep only business context. No private-life details. Store and share records securely.", BODY),
    ("", BODY),
    ("THE PAUSE TEST: Am I sending this because it helps the buyer, or because I'm uncomfortable with the silence?", Font(bold=True, size=11, color=WARM, name="Arial")),
    ("", BODY),
    ("More tools and trust-led outbound for B2B: https://revops.eden.so/", Font(size=10, color=ACCENT, name="Arial", underline="single")),
]
for i, (text, font) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=text)
    c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")

# Trust Graph
tg = wb.create_sheet("Trust Graph")
cols = ["Person", "Company", "Role", "Entry route", "Last buyer signal (their words)", "Signal date",
        "State", "Posture", "Confidence", "Gate", "Ladder rung", "Weakest TU term", "Evidence exchanged",
        "Promises (mine)", "Promise due", "Next move", "Next date / trigger", "Boundary / notes", "Thread link"]
widths = [16, 16, 14, 18, 36, 11, 16, 13, 16, 13, 12, 17, 26, 24, 11, 20, 18, 22, 20]
header(tg, cols, widths)
style_body(tg, len(cols), 200)
for col, vals in {"D": ROUTES, "G": STATES, "H": POSTURES, "I": CONF, "J": GATES, "K": RUNGS, "L": TU, "P": MOVES}.items():
    dropdown(tg, col, vals)
for col in ("F", "O"):
    for r in range(2, 201):
        tg[f"{col}{r}"].number_format = "yyyy-mm-dd"
tg.conditional_formatting.add("A2:S200", FormulaRule(formula=['$G2="Do not contact"'], fill=PatternFill("solid", fgColor="2B2F35"), font=Font(color="FFFFFF")))
tg.conditional_formatting.add("A2:S200", FormulaRule(formula=['AND($O2<>"",$O2<TODAY())'], fill=PatternFill("solid", fgColor="FBEEE8")))
tg.conditional_formatting.add("A2:S200", FormulaRule(formula=['$G2="Active evaluation"'], fill=PatternFill("solid", fgColor=TINT)))
example = ["Maya (illustrative)", "Northline Agency", "Founder", "Content bridge",
           '"Three of the last five kickoffs needed me."', None, "Exploring", "Exploratory", "Directly stated",
           "G4 Safety", "Exploration", "Risk high", "Handoff example discussed; no sample yet",
           "Send written schedule", None, "Clarifier", "Thursday", "Partner approves fees", ""]
for i, v in enumerate(example, start=1):
    tg.cell(row=2, column=i, value=v)
tg.auto_filter.ref = "A1:S200"

# Debugger log
dbg = wb.create_sheet("Debugger Log")
qs = ["Date", "Person", "Exact last buyer signal", "1. What do they believe now?", "2. What don't they believe yet?",
      "3. What are they protecting?", "4. Gate (as a question)", "5. Posture + alternative reading",
      "6. P(outcome) strong/weak/unknown", "7. Relevant value (their words)", "8. Risk or hard constraint",
      "9. Last deposit", "10. Last withdrawal", "11. Missing signal", "12. Smallest move", "13. What NOT to do",
      "Pause Test passed?", "What response would change my diagnosis?"]
header(dbg, qs, [11, 16, 30] + [22] * 13 + [12, 26], fill=INK)
style_body(dbg, len(qs), 100)
dropdown(dbg, "Q", "Yes,No - pause", 100)

# Sprint
sp = wb.create_sheet("30-Day Sprint")
header(sp, ["Day", "Week", "Focus", "Objective", "Done?", "What I observed", "What I first assumed",
            "What I did (incl. justified pause)", "What changed / still unknown", "Metric / evidence", "Tomorrow's adjustment"],
       [6, 11, 24, 30, 8, 28, 26, 28, 26, 22, 24])
plan = [
    (0, "Baseline", "Baseline", "Record last-30-day substantive replies, volunteered context, open promises, dead-end records"),
    (1, "1 Read", "Baseline", "See your actual behavior"), (2, "1 Read", "Fact vs interpretation", "Improve evidence discipline"),
    (3, "1 Read", "Trust Graph Audit", "Preserve relationship context"), (4, "1 Read", "Buyer State Map", "Stop reading keywords as intent"),
    (5, "1 Read", "Gates", "Match moves to open questions"), (6, "1 Read", "TU Diagnostic", "Separate belief, value, risk"),
    (7, "1 Read", "Review", "Learn before adding activity"), (8, "2 Repair", "Missed promises", "Restore reliability"),
    (9, "2 Repair", "Stop conditions", "Prevent unwanted follow-up"), (10, "2 Repair", "Profile", "Reduce identity/relevance doubt"),
    (11, "2 Repair", "Offer explanation", "Make evaluation easier"), (12, "2 Repair", "Trust tax", "Remove unnecessary friction"),
    (13, "2 Repair", "Honest correction", "Repair overreach without a pitch"), (14, "2 Repair", "Coherence", "Align content, profile, DM, offer"),
    (15, "3 Practice", "Mirror", "Make understanding correctable"), (16, "3 Practice", "Clarifier", "Answer an actual uncertainty"),
    (17, "3 Practice", "Insight Gift", "Contribute without obligation"), (18, "3 Practice", "Story", "Explain a mechanism concretely"),
    (19, "3 Practice", "Self-Disclosure", "Make a real limit visible"), (20, "3 Practice", "One trust job", "Reduce message overload"),
    (21, "3 Practice", "TRUST Loop", "Integrate the moves"), (22, "4 Compound", "Content bridge", "Connect public thinking to conversation"),
    (23, "4 Compound", "Modality", "Purposeful format changes"), (24, "4 Compound", "Offer transition", "Offer as continuity"),
    (25, "4 Compound", "Objection diagnosis", "Diagnose before intervening"), (26, "4 Compound", "Rehydration", "Selective, legitimate returns"),
    (27, "4 Compound", "Decision aid", "Proposal someone can decide on"), (28, "4 Compound", "Delivery handoff", "Protect post-sale identity"),
    (29, "Review", "Post-mortems", "Learn without hindsight bias"), (30, "Review", "Next cycle", "Make the system sustainable"),
]
for r, row in enumerate(plan, start=2):
    for c, v in enumerate(row, start=1):
        sp.cell(row=r, column=c, value=v)
style_body(sp, 11, 32)
dropdown(sp, "E", "Done,Justified pause,Skipped", 40)

# Weekly metrics
wm = wb.create_sheet("Weekly Metrics")
mcols = ["Week of", "Cohort / entry route", "Invitations sent", "Accepted", "Unique people messaged", "Unique repliers (incl. refusals)",
         "Refusals", "Substantive conversations", "Qualified opportunities", "Evaluations held", "Proposals sent", "Won", "Lost", "Paused",
         "Signed value", "Cash collected", "Hours spent",
         "Acceptance %", "Reply %", "Substantive %", "Qualified / substantive %", "Proposal win %", "Contact → client %", "Notes: one change next week"]
header(wm, mcols, [11, 18] + [11] * 15 + [11] * 6 + [30])
style_body(wm, len(mcols), 60)


def pct(num, den):
    return f'=IF({den}{{r}}>0,{num}{{r}}/{den}{{r}},"")'


formulas = {"R": pct("D", "C"), "S": pct("F", "E"), "T": pct("H", "E"), "U": pct("I", "H"),
            "V": '=IF((L{r}+M{r})>0,L{r}/(L{r}+M{r}),"")', "W": pct("L", "E")}
for r in range(2, 61):
    for col, f in formulas.items():
        wm[f"{col}{r}"] = f.format(r=r)
        wm[f"{col}{r}"].number_format = "0%"
    wm[f"A{r}"].number_format = "yyyy-mm-dd"
    for col in ("O", "P"):
        wm[f"{col}{r}"].number_format = "#,##0"
wm["X2"] = "Illustrative row: 40 → 20 → 20 → 8 → 5 → 3 → 3 → 2 → 1. Replace with your own."
for col, v in zip("BCDEFGHIJKLMN", ["Example cohort", 40, 20, 20, 8, 1, 5, 3, 3, 2, 1, 1, 0]):
    wm[f"{col}2"] = v

# Post-mortems
pm = wb.create_sheet("Post-Mortems")
header(pm, ["Date", "Person / deal", "Outcome", "Before: what I thought they needed", "Reality: what they signaled",
            "Gate", "Emotion (+ alternative)", "Trust: safer or more pressured?", "Risk increased / decreased",
            "Signal established", "Mistake: where I moved too fast", "Next time", "One supported lesson"],
       [11, 18, 11, 26, 26, 12, 20, 20, 20, 20, 24, 24, 26], fill=WARM)
style_body(pm, 13, 60)
dropdown(pm, "C", "Won,Lost,Stalled,Paused,Progressing", 60)
dropdown(pm, "F", GATES, 60)

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"saved {OUT.relative_to(ROOT)}")
