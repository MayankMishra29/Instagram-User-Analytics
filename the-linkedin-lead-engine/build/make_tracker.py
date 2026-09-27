"""Generate the Content Engine Tracker spreadsheet for The LinkedIn Lead Engine."""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "system" / "Content-Engine-Tracker.xlsx"
SRC = ROOT / "complete-edition"

INK, ACCENT, WARM, TINT = "1D2126", "0F5E5A", "C6613F", "E5F0EE"
HEAD = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY = Font(name="Arial", size=10)
BORDER = Border(bottom=Side(style="thin", color="DFE2E6"))
WRAP = Alignment(wrap_text=True, vertical="top")

AWARE = "Unaware,Problem aware,Solution aware,Provider aware,Ready to buy"
JOBS = "Attraction,Authority,Relationship,Conversion"
ENGINES = "Story,How-to,List,Contrarian,Trend,Lessons learned,Conversation"
CTAS = "Application,Experience,Resource,Conversation/DM,Commercial,None"
STATUS = "Idea,Brief,Drafting,Ready,Published,Reviewed"
OPP = "Conversation,Fit conversation,Qualified opportunity,Proposal,Won,Lost,Paused"
FORMATS = "Text,Image,Document/carousel,Video,Poll"


def header(ws, cols, widths, fill=ACCENT):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HEAD, PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def body(ws, ncols, rows):
    for r in range(2, rows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER


def dropdown(ws, col, values, rows=300):
    dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{rows}")


def md_table_rows(path, min_cols):
    rows = []
    for line in path.read_text().splitlines():
        if line.startswith("|") and not re.match(r"^\|[-\s|]+\|$", line):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) >= min_cols:
                rows.append(cells)
    return rows


def clean(s):
    return re.sub(r"\*\*|\*|✍️ ", "", s).strip()


wb = Workbook()
ws = wb.active
ws.title = "Start here"
ws.column_dimensions["A"].width = 110
lines = [
    ("THE LINKEDIN LEAD ENGINE · CONTENT ENGINE TRACKER", Font(bold=True, size=16, color=INK, name="Arial")),
    ("Part of the Implementation System for the Complete Edition. © Mayank Mishra / ThriveXLabs.", Font(italic=True, size=10, color="4A515B", name="Arial")),
    ("", BODY),
    ("HOW TO USE", Font(bold=True, size=12, color=ACCENT, name="Arial")),
    ("1. Problem Bank: capture buyer phrases after every call and delivery session (Ch 18).", BODY),
    ("2. Post Planner: one row per post. Fill the brief before drafting (Ch 20). Dropdowns keep job, engine, awareness and CTA rung consistent.", BODY),
    ("3. Proof Ledger: every claim you publish needs a source, permission and a limitation (Ch 10).", BODY),
    ("4. Hook Bank: all 50 hook templates from Chapter 22. Mark the ones you've used and what you learned.", BODY),
    ("5. 30-Day Challenge: the daily plan from Chapter 26. A justified reschedule counts as done.", BODY),
    ("6. Weekly Metrics: enter counts; rates calculate with explicit denominators. Read counts next to percentages.", BODY),
    ("7. Opportunities: separate conversations from fit conversations and qualified opportunities. Record sourced vs influenced.", BODY),
    ("", BODY),
    ("RULES THAT KEEP IT HONEST", Font(bold=True, size=12, color=ACCENT, name="Arial")),
    ("• Optimize for qualified attention, not maximum attention.", BODY),
    ("• A download is not a lead. A compliment is not an opportunity.", BODY),
    ("• Never publish an invented experience, client or result. Label illustrations.", BODY),
    ("• Don't change audience, offer, format and CTA at once. Run one experiment per cycle.", BODY),
    ("", BODY),
    ("Companion book and more tools: https://revops.eden.so/", Font(size=10, color=ACCENT, name="Arial", underline="single")),
]
for i, (t, f) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True)

pb = wb.create_sheet("Problem Bank")
header(pb, ["Date", "Buyer phrase (exact words)", "Source", "OK to publish?", "Situation", "Decision they face",
            "Pillar", "Evidence available", "Possible next step", "Post ideas (by job)"], [11, 34, 16, 12, 28, 28, 18, 24, 22, 34])
body(pb, 10, 150)
dropdown(pb, "D", "Yes,Anonymized only,No")
pb["B2"], pb["C2"], pb["D2"], pb["E2"], pb["F2"] = ('"We get inquiries, but most want something cheaper." (illustrative)',
                                                  "Discovery call", "Anonymized only", "Specialist agency attracting early-stage businesses",
                                                  "Change targeting, offer or qualifying info?")

pp = wb.create_sheet("Post Planner")
cols = ["Post ID", "Publish date", "Status", "Reader & situation", "Awareness", "Primary job", "Engine", "Format",
        "Thesis (one sentence)", "Hook", "Rehook", "Evidence / proof", "Payoff", "Boundary", "CTA rung", "CTA text & destination",
        "Impressions", "Relevant comments", "Substantive DMs", "Fit conversations", "Opportunities", "Next question this post revealed"]
header(pp, cols, [9, 11, 11, 26, 14, 13, 14, 13, 30, 30, 30, 24, 24, 20, 14, 26, 11, 11, 11, 11, 11, 30])
body(pp, len(cols), 200)
for col, v in {"C": STATUS, "E": AWARE, "F": JOBS, "G": ENGINES, "H": FORMATS, "O": CTAS}.items():
    dropdown(pp, col, v)
for r in range(2, 201):
    pp[f"B{r}"].number_format = "yyyy-mm-dd"
pp.conditional_formatting.add("A2:V200", FormulaRule(formula=['$C2="Published"'], fill=PatternFill("solid", fgColor=TINT)))

pl = wb.create_sheet("Proof Ledger")
header(pl, ["Claim", "Source / record", "Date", "Permission", "Limitations", "Approved wording", "Used in post IDs"],
       [34, 26, 11, 14, 30, 36, 16], fill=INK)
body(pl, 7, 100)
dropdown(pl, "D", "Own data,Client approved,Public source,Pending,Not allowed")

hb = wb.create_sheet("Hook Bank")
header(hb, ["#", "Category", "Template", "Example", "Used? (post ID)", "What I learned"], [5, 16, 48, 52, 14, 30], fill=WARM)
cat = ""
r = 2
for line in (SRC / "06-part6-hooks.md").read_text().splitlines():
    m = re.search(r'<div class="table-caption">(.*?)</div>', line)
    if m:
        cat = m.group(1)
    if "Chapter 23" in line:
        break
    if re.match(r"^\| \d\d \|", line):
        n, tpl, ex = [c.strip() for c in line.strip().strip("|").split("|")][:3]
        for i, v in enumerate([n, cat, clean(tpl), clean(ex)], start=1):
            hb.cell(row=r, column=i, value=v)
        r += 1
body(hb, 6, r)
print("hooks loaded:", r - 2)

ch = wb.create_sheet("30-Day Challenge")
header(ch, ["Day", "Task", "Output", "Done?", "Evidence created", "Reader response", "Time spent (min)", "One adjustment"],
       [6, 60, 30, 12, 26, 26, 12, 26])
rows = [x for x in md_table_rows(SRC / "08-part6-swipe-challenge.md", 3) if re.match(r"^\*\*\d+\*\*$", x[0])]
for i, (d, task, out) in enumerate([(clean(a), clean(b), clean(c)) for a, b, c in [x[:3] for x in rows]], start=2):
    ch.cell(row=i, column=1, value=int(d))
    ch.cell(row=i, column=2, value=task)
    ch.cell(row=i, column=3, value=out)
body(ch, 8, 32)
dropdown(ch, "D", "Done,Rescheduled (reason noted),Skipped", 40)
print("challenge days:", len(rows))

wm = wb.create_sheet("Weekly Metrics")
mcols = ["Week of", "Posts published", "Impressions", "Members reached", "Relevant comments", "Substantive new conversations",
         "Fit conversations", "Qualified opportunities", "Proposals sent", "Won", "Lost", "Open", "Resource requests",
         "Resource used (observed)", "Hours on content + replies", "Cash collected",
         "Fit rate (fit ÷ substantive)", "Opportunity rate (qualified ÷ fit)", "Win rate (won ÷ decided)", "Resource use rate", "Notes: one change next week"]
header(wm, mcols, [11] + [12] * 15 + [14, 14, 14, 13, 30])
body(wm, len(mcols), 60)
for r in range(2, 61):
    wm[f"A{r}"].number_format = "yyyy-mm-dd"
    wm[f"Q{r}"] = f'=IF(F{r}>0,G{r}/F{r},"")'
    wm[f"R{r}"] = f'=IF(G{r}>0,H{r}/G{r},"")'
    wm[f"S{r}"] = f'=IF((J{r}+K{r})>0,J{r}/(J{r}+K{r}),"")'
    wm[f"T{r}"] = f'=IF(M{r}>0,N{r}/M{r},"")'
    for c in "QRST":
        wm[f"{c}{r}"].number_format = "0%"
for c, v in zip("BCDEFGHIJKLMN", [12, None, None, None, 12, 6, 3, 2, 1, 0, 1, None, None]):
    if v is not None:
        wm[f"{c}2"] = v
wm["U2"] = "Illustrative row from Chapter 21 (invented, not a benchmark). Replace with your own."

op = wb.create_sheet("Opportunities")
header(op, ["Person / company", "Origin (post ID, resource, referral)", "Sourced or influenced", "Problem in their words",
            "Fit evidence", "Open questions", "Decision path & timing", "Next action · owner · date", "Status", "Boundary", "Value (signed)", "Cash collected"],
       [20, 22, 13, 30, 26, 24, 22, 26, 16, 18, 12, 12], fill=WARM)
body(op, 12, 150)
dropdown(op, "C", "Sourced,Influenced,Both,Unknown")
dropdown(op, "I", OPP)

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("saved", OUT.relative_to(ROOT))
